"""M10: importing spreadsheets, lab exports (NED-Net, TRACER, QuantStudio) and Prism files."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

SAMPLES = Path(__file__).resolve().parents[1] / "sample_data"


def _need(name: str) -> Path:
    p = SAMPLES / name
    if not p.exists():
        pytest.skip(f"{name} is not in sample_data/ (local lab files are not in the repository)")
    return p


# ---------------------------------------------------------------------------
# Reading spreadsheets
# ---------------------------------------------------------------------------

def test_csv_with_semicolons_and_decimal_commas(tmp_path):
    from asterism.io.sheets import guess_header_row, read_file
    f = tmp_path / "data.csv"
    f.write_text("Run on 2026-10-01\n\nGroup;Value\nControl;1,5\nDrug;2,25\n", encoding="utf-8")
    rows = read_file(f)["data"]
    assert guess_header_row(rows) == 2
    assert rows[3] == ["Control", 1.5] and rows[4][1] == 2.25


def test_xlsx_round_trip(tmp_path):
    from openpyxl import Workbook
    from asterism.io.sheets import read_file
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet A"
    ws.append(["Day", "A", "B"])
    ws.append([dt.datetime(2026, 10, 1), 1, "2"])
    f = tmp_path / "x.xlsx"
    wb.save(f)
    rows = read_file(f)["Sheet A"]
    assert rows[1][0] == dt.datetime(2026, 10, 1) and rows[1][2] == 2.0


# ---------------------------------------------------------------------------
# Building tables
# ---------------------------------------------------------------------------

def test_wide_column_and_grouped_with_replicates():
    from asterism.io.build import Plan, build_table, guess_plan
    rows = [["Animal", "Control", "Drug"], ["m1", 1, 3], ["m2", 2, "n/a"], ["m3", None, 5]]
    plan = guess_plan(rows)
    assert plan.layout == "wide" and plan.roles[1] == plan.roles[2] == "set"
    t, notes = build_table(rows, Plan("column", "wide", 0, {1: "set", 2: "set"}))
    assert t.titles[:2] == ["Control", "Drug"] and list(t.values(1)) == [3, 5]
    assert "1 text cell" in notes[0]
    rows = [["Day", "WT 1", "WT 2", "KO 1", "KO 2"], ["Day 1", 1, 2, 3, 4], ["Day 2", 5, 6, 7, 8]]
    t, _ = build_table(rows, Plan("grouped", "wide", 0, {0: "row", 1: "set", 2: "set", 3: "set", 4: "set"}))
    assert t.y_format == "replicates" and t.width() == 2 and t.titles[:2] == ["WT", "KO"]
    assert t.row_label(1) == "Day 2" and t.get(1, t.flat(1, 1)) == 8


def test_wide_xy_with_dates():
    from asterism.doc.dates import fmt
    from asterism.io.build import Plan, build_table
    rows = [["Date", "Weight"], [dt.datetime(2026, 1, 1), 20.0], [dt.datetime(2026, 1, 8), 21.5]]
    t, _ = build_table(rows, Plan("xy", "wide", 0, {0: "x", 1: "set"}))
    assert t.x_format == "date" and t.has_x()
    assert fmt(t.x_value(1)) == "2026-01-08" and t.x_value(1) - t.x_value(0) == 7


def test_long_layouts_for_every_table_type():
    from asterism.io.build import Plan, build_table
    rows = [["mouse", "group", "week", "value"]]
    for m, g in (("a", "WT"), ("b", "WT"), ("c", "KO")):
        for w in (1, 2):
            rows.append([m, g, w, (ord(m) - 96) * 10 + w])
    roles = {0: "subject", 1: "group", 2: "factor2", 3: "value"}
    # grouped, repeated measures: each mouse keeps its sub-column across weeks
    t, _ = build_table(rows, Plan("grouped", "long", 0, roles))
    assert t.titles[:2] == ["WT", "KO"] and t.sub_titles[0][:2] == ["a", "b"]
    assert [t.row_label(i) for i in range(2)] == ["1", "2"]
    assert t.get(1, t.flat(0, 1)) == 22 and t.get(0, t.flat(1, 0)) == 31
    # nested: mice side by side, their values below
    t, _ = build_table(rows, Plan("nested", "long", 0, {0: "subject", 1: "group", 3: "value"}))
    assert t.table_type == "nested" and t.sub_titles[0][:2] == ["a", "b"] and t.get(1, t.flat(0, 0)) == 12
    # column, averaged per subject
    t, notes = build_table(rows, Plan("column", "long", 0, {0: "subject", 1: "group", 3: "value"},
                                      average_subjects=True))
    assert sorted(t.values(0)) == [11.5, 21.5] and "averaged" in notes[0]
    # xy: X from the second factor, numeric and sorted
    t, _ = build_table(rows, Plan("xy", "long", 0, {0: "subject", 1: "group", 2: "x", 3: "value"}))
    assert t.x_value(0) == 1 and t.x_value(1) == 2
    # contingency: counts of combinations
    crow = [["outcome", "treat"], ["alive", "A"], ["dead", "A"], ["alive", "B"], ["alive", "A"]]
    t, notes = build_table(crow, Plan("contingency", "long", 0, {0: "group", 1: "factor2"}))
    assert t.titles[:2] == ["alive", "dead"] and t.get(0, t.flat(0)) == 2 and "Counted" in notes[0]
    # survival: time, event, group
    srow = [["days", "died", "arm"], [5, 1, "A"], [9, 0, "A"], [3, 1, "B"]]
    t, _ = build_table(srow, Plan("survival", "long", 0, {0: "x", 1: "event", 2: "group"}))
    assert t.x_value(1) == 9 and t.get(1, t.flat(0)) == 0 and t.get(2, t.flat(1)) == 1


# ---------------------------------------------------------------------------
# Lab exports
# ---------------------------------------------------------------------------

def test_nednet_export_becomes_typed_tables():
    from asterism.io import labs
    from asterism.io.sheets import read_file
    sh = read_file(_need("seizures_results_graph_data_ALLweeks.xlsx"))
    src, items = labs.detect(sh)
    assert src == "NED-Net" and len(items) == 11
    rate = next(i for i in items if i.metric == "rate_per_h")
    (t,), notes = labs.build(rate, sh)
    assert t.table_type == "column" and t.titles[:2] == ["SV2A", "Control"]
    assert len(t.values(0)) + len(t.values(1)) == 21 and "355676" in notes[0] and ".0" not in notes[0]
    daily = next(i for i in items if i.sheet == "Daily_counts_byGroup")
    (t,), _ = labs.build(daily, sh)
    assert t.table_type == "grouped" and t.sub_titles[0][:3] == ["25", "31", "33"] and t.row_label(0) == "1"
    daily.layout = "xy_x"
    (t,), _ = labs.build(daily, sh)
    assert t.table_type == "xy" and t.x_value(0) == 1


def test_tracer_cells_become_nested_tables(tmp_path):
    from openpyxl import Workbook
    from asterism.io import labs
    from asterism.io.sheets import read_file
    wb = Workbook()
    ws = wb.active
    ws.title = "Cells (long)"
    ws.append(["cell_id", "animal_id", "group", "file_name", "metric", "value"])
    for cell, animal, group, v in (("c1", "m1", "WT", 10), ("c2", "m1", "WT", 12), ("c3", "m2", "WT", 9),
                                   ("c4", "m3", "KO", 20)):
        ws.append([cell, animal, group, "f.abf", "amplitude", v])
        ws.append([cell, animal, group, "f.abf", "frequency", v / 10])
    f = tmp_path / "cohort_summary.xlsx"
    wb.save(f)
    sh = read_file(f)
    src, items = labs.detect(sh)
    assert src == "TRACER" and [i.metric for i in items] == ["amplitude", "frequency"]
    (t,), _ = labs.build(items[0], sh)
    assert t.table_type == "nested" and t.titles[:2] == ["WT", "KO"] and t.sub_titles[0][:2] == ["m1", "m2"]
    assert [t.get(r, t.flat(0, 0)) for r in range(2)] == [10, 12]
    items[0].layout = "per_subject"
    (t,), _ = labs.build(items[0], sh)
    assert sorted(t.values(0)) == [9, 11]


def _quantstudio_csv(tmp_path) -> Path:
    lines = ["* Block Type = 96-Well Block", "* Chemistry = TAQMAN", "",
             "Well,Well Position,Omit,Sample Name,Target Name,Task,Reporter,CT"]
    data = [("WT 1", 20.0, 25.0), ("WT 2", 20.5, 25.0), ("KO 1", 20.0, 27.0), ("KO 2", 21.0, 28.5)]
    w = 1
    for smp, ref, tgt in data:
        for target, ct in (("Gapdh", ref), ("Syt1", tgt)):
            for k in range(2):                          # technical duplicates
                lines.append(f"{w},A{w},false,{smp},{target},UNKNOWN,FAM,{ct + (0.1 if k else -0.1)}")
                w += 1
    lines.append(f"{w},H1,false,water,Syt1,NTC,FAM,Undetermined")
    lines.append(f"{w + 1},H2,true,WT 1,Syt1,UNKNOWN,FAM,40")
    lines += ["", "Endogenous Control,Gapdh", "Reference Sample,WT 1"]
    f = tmp_path / "run.csv"
    f.write_text("\n".join(lines))
    return f


def test_quantstudio_relative_expression(tmp_path):
    from asterism.io import labs, quantstudio as Q
    from asterism.io.sheets import read_file
    sh = read_file(_quantstudio_csv(tmp_path))
    d = Q.defaults(sh)
    assert d["reference_genes"] == ["Gapdh"] and d["group_list"] == ["WT", "KO"] and d["control_group"] == "WT"
    res = Q.compute(sh, {})
    # ΔCt: WT 1 = 5, WT 2 = 4.5 → control mean 4.75; KO 1 = 7 → ΔΔCt 2.25
    assert res["dct"][("KO 1", "Syt1")] == pytest.approx(7.0)
    assert res["rel"][("KO 1", "Syt1")] == pytest.approx(2 ** -2.25)
    notes = " ".join(res["notes"])
    assert "1 no-template" in notes and "1 omitted" in notes
    by_sample = Q.compute(sh, {"relative_to": "sample", "reference_sample": "WT 1"})
    assert by_sample["rel"][("WT 1", "Syt1")] == pytest.approx(1.0)
    src, items = labs.detect(sh)
    assert src == "QuantStudio"
    (t,), _ = labs.build(items[0], sh, {})
    assert t.name == "Syt1 relative expression" and t.titles[:2] == ["WT", "KO"]
    (t,), _ = labs.build(next(i for i in items if i.layout == "qpcr_ct"), sh, {})
    assert t.table_type == "grouped" and t.row_label(0) == "WT 1"


def test_quantstudio_real_export_matches_its_rq():
    """Our ΔΔCt, relative to the run's own reference sample, equals QuantStudio's RQ column."""
    from asterism.io import quantstudio as Q
    from asterism.io.sheets import read_file
    sh = read_file(_need("Syt1 primary gRNA3 200708.xls"))
    foot = Q.footer(sh)
    assert foot == {"reference_gene": "B-act", "reference_sample": "Control 3"}
    res = Q.compute(sh, {"relative_to": "sample", "reference_sample": foot["reference_sample"]})
    name, hr = Q.find_results(sh)
    rows = sh[name]
    h = rows[hr]
    i_rq, i_s, i_t = h.index("RQ"), h.index("Sample Name"), h.index("Target Name")
    rq = {(r[i_s], r[i_t]): r[i_rq] for r in rows[hr + 1:] if len(r) > i_rq and isinstance(r[i_rq], float)}
    assert res["rel"] and all(res["rel"][k] == pytest.approx(rq[k], rel=1e-6) for k in res["rel"])


# ---------------------------------------------------------------------------
# Prism files
# ---------------------------------------------------------------------------

def test_prism_file_tables_exclusions_and_analyses():
    from asterism.analysis import engine
    from asterism.doc.model import Project
    from asterism.io.prism import into_project, read
    imp = read(_need("Seizure_data.prism"))
    assert len(imp.tables) == 37 and len(imp.analyses) == 26 and len(imp.infos) == 1
    assert all(isinstance(x, str) for t in imp.tables for x in t.titles)
    assert sum(len(t.excluded) for t in imp.tables) == 204
    assert not any("graph" in n for n in imp.notes) and any("layout" in n for n in imp.notes)
    p = Project(engine)
    into_project(imp, p)
    assert not [a for a in p.analyses if a.result is None or a.result.error]
    # Prism's graphs: one ASTERISM graph each (titles, ranges, intervals, log scale), plus a default
    # graph for tables Prism had none for
    assert len(imp.graphs) == 37 and len(p.graphs) == 39
    log = next(g for g in p.graphs if g.name == "Seizures/day log10")
    assert log.spec.y.scale == "log10" and log.spec.y.vmin == pytest.approx(0.01) and log.spec.y.title == "Seizures/day"
    lin = next(g for g in p.graphs if g.name == "% change from week 1 (old)")
    assert (lin.spec.y.vmin, lin.spec.y.vmax, lin.spec.y.interval, lin.spec.y.auto) == (-500, 2500, 500, False)
    daily = next(g for g in p.graphs if g.name == "Daily seizures")
    assert daily.spec.y.title == "Events/animal-hour"
    corr = next(a for a in p.analyses if a.analysis_type == "correlation")
    rows = {r[0]: r[1] for r in corr.result.pages[0].rows if r[1]}
    assert rows["Spearman r"] == ["-0.7714", "0.6845"]          # Prism: −0.7714, 0.6845


def test_pzfx(tmp_path):
    from asterism.io.prism import read
    f = tmp_path / "t.pzfx"
    f.write_text("""<?xml version="1.0" encoding="UTF-8"?>
<GraphPadPrismFile xmlns="http://graphpad.com/prism/Prism.htm" PrismXMLVersion="5.00">
<InfoSequence><Ref ID="Info0"/></InfoSequence>
<Info ID="Info0"><Title>Project info 1</Title><Notes>Mice from batch 3</Notes>
<Constant><Name>Experimenter</Name><Value>ML</Value></Constant></Info>
<TableSequence><Ref ID="Table0"/></TableSequence>
<Table ID="Table0" XFormat="none" TableType="TwoWay" YFormat="replicates" Replicates="2">
<Title>Seizures</Title>
<RowTitlesColumn Width="80"><Subcolumn><d>Week 1</d><d>Week 2</d></Subcolumn></RowTitlesColumn>
<YColumn Width="100" Decimals="1" Subcolumns="2"><Title>WT</Title>
<Subcolumn><d>1</d><d Excluded="1">2</d></Subcolumn><Subcolumn><d>3</d><d>4</d></Subcolumn></YColumn>
<YColumn Width="100" Decimals="1" Subcolumns="2"><Title>KO</Title>
<Subcolumn><d>5</d><d>6</d></Subcolumn><Subcolumn><d>7</d><d/></Subcolumn></YColumn>
</Table>
</GraphPadPrismFile>""")
    imp = read(f)
    (t,) = imp.tables
    assert t.table_type == "grouped" and t.width() == 2 and t.titles[:2] == ["WT", "KO"]
    assert t.row_label(1) == "Week 2" and t.is_excluded(1, t.flat(0, 0)) and t.get(1, t.flat(1, 1)) is None
    assert imp.infos[0].notes == "Mice from batch 3" and imp.infos[0].fields["Experimenter"] == "ML"


# ---------------------------------------------------------------------------
# Dates
# ---------------------------------------------------------------------------

def test_dates_parse_format_and_axis():
    from asterism.app.graph.xy_layout import DateXMap, x_range
    from asterism.doc import dates
    from asterism.doc.spec import GraphSpec
    d = dates.parse("2026-10-26")
    assert dates.fmt(d) == "2026-10-26" and dates.parse("26/10/2026") == d
    spec = GraphSpec()
    r = x_range(spec, [d, d + 30], date=True)
    assert r.interval in (7.0, 14.0) and r.vmin <= d and r.vmax >= d + 30
    m = DateXMap(spec.x, r, 100.0, d)
    assert m.label(d) == "26 Oct"
    spec.x.date_fmt = "days"
    assert DateXMap(spec.x, r, 100.0, d).label(d + 7) == "7"


# ---------------------------------------------------------------------------
# QuantStudio standard curves (absolute quantification)
# ---------------------------------------------------------------------------

def _std_curve_csv(tmp_path) -> Path:
    import math
    lines = ["# Instrument Type : QuantStudio 5 System", "",
             '"Well","Sample Name","Target Name","Task","Cq","Quantity","Omit",']
    w = 0
    for q in (1e8, 1e7, 1e6, 1e5):
        for k in range(2):
            w += 1
            lines.append(f'"A{w}","No Sample","ITR","STANDARD","{40 - 3.32 * math.log10(q) + (0.05 if k else -0.05):.3f}",'
                         f'"{q}","false",')
    for smp, cq in (("Batch 1", 16.72), ("Batch 2", 20.04)):
        for k in range(3):
            w += 1
            lines.append(f'"B{w}","{smp}","ITR","UNKNOWN","{cq}","-","false",')
    lines.append(f'"C1","water","ITR","NTC","Undetermined","-","false",')
    folder = tmp_path / "run.eds"
    folder.mkdir()
    (folder / "Results.csv").write_text("\n".join(lines))
    (folder / "Plate QC.csv").write_text('"Name","FLAGS"\n"run.eds","0"\n')
    return folder


def test_standard_curve_and_titres(tmp_path):
    from asterism.io import labs, quantstudio as Q
    from asterism.io.sheets import read_files
    sh = read_files([_std_curve_csv(tmp_path)])                    # a folder of separate CSV files
    assert set(sh) == {"Results", "Plate QC"} and Q.has_standards(sh)
    cv = Q.standard_curves(sh)["ITR"]
    assert cv["slope"] == pytest.approx(-3.32, abs=1e-6) and cv["intercept"] == pytest.approx(40.0, abs=1e-6)
    assert cv["efficiency"] == pytest.approx((10 ** (1 / 3.32) - 1) * 100) and cv["r2"] > 0.999
    res = Q.quantify(sh)
    assert res["quantities"][("Batch 1", "ITR")][0] == pytest.approx(10 ** ((16.72 - 40) / -3.32))
    # titre: 1:100 dilution, 2 µL per reaction, per mL, single-stranded genomes
    t2 = Q.quantify(sh, {"dilution": 100, "volume_ul": 2, "per": "ml", "single_stranded": True})
    assert t2["quantities"][("Batch 1", "ITR")][0] == pytest.approx(
        res["quantities"][("Batch 1", "ITR")][0] * 100 * 2 / 2 * 1000)
    assert t2["unit"] == "copies per mL"
    src, items = labs.detect(sh)
    assert [i.layout for i in items if i.selected] == ["qpcr_std", "qpcr_qty"]
    (std,), notes = labs.build(items[0], sh, {})
    assert std.table_type == "xy" and std.x_value(0) == pytest.approx(8.0) and std.import_analyses[0] == ("linreg", [0], {})
    # the unknowns, on the curve: one data set per sample, a cross per well
    assert std.titles[:3] == ["ITR standards", "Batch 1", "Batch 2"]
    assert std.dataset_styles[1]["symbol_shape"] == "cross" and len(std.values(1)) == 3
    r = next(i for i in range(std.nrows_used()) if std.get(i, std.flat(1, 0)) is not None)
    assert std.x_value(r) == pytest.approx((16.72 - 40) / -3.32) and std.get(r, std.flat(1, 0)) == 16.72
    (qty,), _ = labs.build(items[1], sh, {"standard": {"dilution": 100}})
    assert qty.table_type == "grouped" and qty.row_label(1) == "Batch 2" and qty.width() == 3


def test_real_standard_curve_matches_quantstudio():
    from asterism.io import quantstudio as Q
    from asterism.io.sheets import read_files
    folder = _need("2018-02-27_122901 std curve")
    sh = read_files([folder])
    cv = Q.standard_curves(sh)
    # QuantStudio (from unrounded Cq): CMV slope −3.3696, intercept 35.8058, efficiency 98.0483 %
    assert cv["CMV"]["slope"] == pytest.approx(-3.3696, abs=5e-4)
    assert cv["CMV"]["intercept"] == pytest.approx(35.8058, abs=5e-3)
    assert cv["ITR"]["efficiency"] == pytest.approx(92.292274, abs=0.05)
    q = Q.quantify(sh)["quantities"]
    assert q[("#175 Vale 1", "CMV")][0] == pytest.approx(1.0557656e7, rel=2e-3)



def test_prism_graph_binary_structure():
    """The record structure of Prism's graph files (containers with end markers, plain records)."""
    import struct
    from asterism.io.prism_graphs import decode, parse
    rec = lambda tag, pl=b"": struct.pack("<HI", tag, len(pl)) + pl
    end = lambda tag: struct.pack("<H", (tag & 0x3fff) | 0x4000)
    axis = lambda lo, hi: struct.pack("<5d", 0, 0, lo, hi, lo) + b"\0" * 20
    seg = lambda iv: b"\0" * 30 + struct.pack("<H", 1200) + b"\0" * 22 + struct.pack("<d", iv) + b"\0" * 72
    text = lambda kind, box, s, vertical, role=b"\x30\x03": (rec(0x8020, struct.pack("<H4h", kind, *box) + b"\0" * 18 + role)
                                           + rec(0x0071, b"\x34\0\0\0" + (b"\1" if vertical else b"\0") + b"\0\0\0")
                                           + rec(0x802b, b"\0" * 4) + rec(0x002c, s.encode()) + end(0x802b) + end(0x8020))

    def dataset(lo, hi, bars, shape, fill, dash=0):
        st = bytearray(317)
        st[15] = 0x10 if bars else 0
        st[28], st[34], st[81] = 3, shape, 3
        st[208:212] = b"\xff" + bytes.fromhex(fill)
        st[224:228] = b"\xff\0\0\0"
        rng = struct.pack("<6d", 0, 0, 0, 0, lo, hi)[:16] + b"\0" * 10 + struct.pack("<4d", 0, 0, lo, hi)
        rng += b"\0" * 8 + bytes([dash]) + b"\0" * 20                     # byte 66: line pattern
        return (rec(0x8020, b"\0" * 30) + rec(0x8028, rng) + rec(0x009c, bytes(st)) + end(0x8028) + end(0x8020))
    body = (rec(0x8012, b"3\0") + rec(0x8012, b"x\0")                    # a plain record with the 0x8000 bit
            + rec(0x0017, axis(1, 3)) + rec(0x0017, axis(-2, 3)) + rec(0x0017, axis(0, 100))
            + rec(0x00a5, b"\0" * 64 + b"\0" * 4) + rec(0x00a5, b"\0" * 64 + b"\1\0\1\0") + rec(0x00a5, b"\0" * 68)
            + b"".join(rec(0x8025, seg(iv)) + end(0x8025) for iv in (1, 1, 1, 50, 1, 1, 1, 1, 1))
            + text(2, (100, 50, 300, 80), "My graph", False, b"\x33\x03") + text(2, (100, 900, 300, 940), "Day", False)
            + text(2, (10, 300, 40, 700), "Cells", True) + text(2, (10, 300, 40, 700), "Y1Title", True)
            + dataset(5.0, 9.0, bars=True, shape=10, fill="999999") + dataset(1.0, 2.0, bars=False, shape=4, fill="ffffff", dash=1)
            + end(0x8012))
    buf = b"PCFFGRA4" + rec(0x8001, b"\0" * 12) + body + end(0x8001)
    nodes, stop = parse(buf)
    assert stop == len(buf) and nodes[0].tag == 0x8001
    gf = decode(buf)
    assert (gf.title, gf.x_title, gf.y_title) == ("My graph", "Day", "Cells")
    assert gf.axes["y"].log and gf.axes["y"].vmin == pytest.approx(0.01) and gf.axes["y"].vmax == pytest.approx(1000)
    assert gf.axes["x"].vmin == 1 and gf.axes["x"].vmax == 3 and gf.axes["x"].interval == 1
    assert not gf.title_visible                         # role bit 0x02: hidden in Prism
    assert gf.axes["y"].length_mm == pytest.approx(50.8)   # 1200/600 inch
    a, b = gf.sets
    assert (a.ymin, a.ymax, a.bars, a.shape, a.fill, a.error) == (5.0, 9.0, True, 10, "#999999", 3)
    assert (b.ymin, b.ymax, b.bars, b.shape, b.dash, a.dash) == (1.0, 2.0, False, 4, 1, 0)


def test_prism_graph_styles_onto_spec():
    """Decoded styles land on the matching data sets of an ASTERISM graph."""
    from asterism.analysis import engine
    from asterism.doc.model import Project
    from asterism.io.prism_graphs import AxisFormat, GraphFormat, SetStyle, apply
    p = Project(engine)
    t = p.new_table("T", "column")
    for c, vals in ((0, [1, 2, 3]), (1, [10, 20, 30])):
        for r, v in enumerate(vals):
            t.set(r, c, v)
    g = p.new_graph(t)
    gf = GraphFormat(title="T", title_visible=False, legend=True,
                     axes={"y": AxisFormat(0, 40, 10, False, 45.0), "x": AxisFormat(length_mm=30.0)},
                     sets=[SetStyle(ymin=10, ymax=30, bars=False, shape=10, symbol_border="#123456", error=2),
                           SetStyle(ymin=1, ymax=3, bars=True, fill="#999999", shape=5)])
    apply(gf, g.spec, "column", t)
    s0, s1 = g.spec.style(0), g.spec.style(1)
    assert s0.appearance == "bar" and s0.bar_fill == "#999999" and s0.symbol_shape == "circle"
    assert s1.appearance == "scatter" and s1.symbol_shape == "open_circle" and s1.symbol_color == "#123456"
    assert s1.line_at == "mean_sd"
    assert not g.spec.show_title and g.spec.show_legend
    assert g.spec.y_len_mm == 45.0 and g.spec.col_width_mm == 15.0


def test_real_prism_dashed_line():
    """A solid/dashed pair saved in Prism: data set B dashed, empty X plotted by row number."""
    from asterism.io.prism import into_project, read
    from asterism.analysis import engine
    from asterism.doc.model import Project
    from asterism.app.graph.xy_layout import row_x
    path = _need("dash_dashed.prism")
    p = Project(engine)
    into_project(read(str(path)), p)
    g = p.graphs[0]
    assert [g.spec.style(c).line_dash for c in (0, 1)] == ["solid", "dash"]
    assert row_x(g.table) == [1.0, 2.0, 3.0, 4.0, 5.0]


def test_prism_one_way_anovas_match_prism():
    """One-way ANOVAs read from a .prism file (ordinary with Tukey or Dunnett, repeated
    measures with Geisser-Greenhouse and Šídák, Kruskal-Wallis with Dunn, Brown-Forsythe and
    Welch with Dunnett's T3) and recomputed: the same numbers as Prism's own results."""
    import json
    import zipfile
    from asterism.analysis import engine
    from asterism.doc.model import Project
    from asterism.io.prism import into_project, read
    path = _need("Anovas et al.prism")
    imp = read(str(path))
    assert not any("one way" in n.lower() for n in imp.notes)
    p = Project(engine)
    into_project(imp, p)
    # Prism's adjusted P values, per analysis title + family
    z = zipfile.ZipFile(path)
    want = {}
    for n in z.namelist():
        if n.endswith("/results.json"):
            base = n.rsplit("/", 1)[0]
            sh = json.loads(z.read(base + "/sheet.json"))
            par = json.loads(z.read(base + "/parameters.json"))["content"]
            fam = par["postTests"]["pairwiseTests"]["pairwiseDesign"]["type"]
            ps = []

            def walk(o):
                if isinstance(o, dict):
                    if "pValueAdjusted" in o:
                        ps.append(o["pValueAdjusted"]["value"])
                    for v in o.values():
                        walk(v)
            walk(json.loads(z.read(n))["content"])
            want[(sh["title"].split(" of ")[0], fam)] = sorted(ps)
    got = {}
    for a in p.analyses:
        fam = "CONTROL_VS_ALL" if a.params.get("compare") == "control" else "ALL_VS_ALL"
        comps = [float(vals[-1]) for pg in a.result.pages for lab, vals, _ in pg.rows if " vs. " in lab]
        got[(a.name.split(" of ")[0], fam)] = sorted(comps)
    pairs = {("One-way ANOVA", "ALL_VS_ALL"): ("Ordinary one-way ANOVA", "ALL_VS_ALL"),
             ("One-way ANOVA", "CONTROL_VS_ALL"): ("Ordinary one-way ANOVA", "CONTROL_VS_ALL"),
             ("Repeated-measures one-way ANOVA", "ALL_VS_ALL"): ("RM one-way ANOVA", "ALL_VS_ALL"),
             ("Welch and Brown-Forsythe ANOVA", "ALL_VS_ALL"): ("Brown-Forsythe and Welch ANOVA tests",
                                                               "ALL_VS_ALL"),
             ("Kruskal-Wallis test", "ALL_VS_ALL"): ("Kruskal-Wallis test", "ALL_VS_ALL")}
    for ours, theirs in pairs.items():
        assert got[ours] == pytest.approx(want[theirs], abs=6e-5), ours     # results show 4 decimals
    # tables: sub-column titles and date X
    nested = next(t for t in p.tables if t.table_type == "nested")
    assert nested.sub_titles[0] == ["Mouse 1", "Mouse 2", "Mouse 3"]
    xy = next(t for t in p.tables if t.table_type == "xy")
    assert xy.x_format == "date" and xy.x_value(0) == 45658.0          # 2025-01-01 as a day number
