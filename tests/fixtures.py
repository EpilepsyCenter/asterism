"""Synthetic workbook builders that reproduce the frozen apps' export writers.

These mirror TRACER `cohort_export.py` and NED-Net `results.py` exactly (same headers,
same ', '.join tag encoding, same two-row-header layout), so importing them exercises
the real paths without needing a real file on disk. A real export is the final
acceptance check (SPEC §13 step 2), not a prerequisite.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def build_tracer_cohort_summary(path: Path) -> Path:
    """cohort_summary.xlsx: Cohort (meta preamble), Cells (long), Cells (wide),
    and one per-event sheet — the sheets the importer cares about."""
    wb = Workbook()

    # Cohort meta sheet (key/value preamble; classifier should ignore it).
    ws = wb.active
    ws.title = "Cohort"
    ws.append(["Cohort export"])
    ws.append(["Generated", "2026-08-06T00:00:00Z"])
    ws.append(["Analysis type", "mEPSC"])

    # Cells (long): [cell_id, animal_id, group, file_name, metric, value]
    ws = wb.create_sheet("Cells (long)")
    ws.append(["cell_id", "animal_id", "group", "file_name", "metric", "value"])
    for animal, grp in (("m1", "WT"), ("m2", "WT"), ("m3", "KO"), ("m4", "KO")):
        for cell in range(3):
            base = 10.0 if grp == "WT" else 16.0
            ws.append([f"{animal}c{cell}", animal, grp, f"{animal}.dat",
                       "amplitude_pA", base + cell])

    # Cells (wide): prefix + scalar metrics + meta.* ; tags ', '.join encoded.
    ws = wb.create_sheet("Cells (wide)")
    ws.append(["group", "file_name", "file_path", "cell_id", "animal_id",
               "series_key (0-indexed)", "series_specific_tags", "group_tags",
               "amplitude_pA", "frequency_Hz", "meta.rig"])
    wide_rows = [
        ("WT", "m1.dat", "/x/m1.dat", "m1c0", "m1", "0:3", "male", "WT, adult", 10.0, 2.0, "rigA"),
        ("WT", "m2.dat", "/x/m2.dat", "m2c0", "m2", "0:3", "female", "WT, adult", 11.0, 2.1, "rigA"),
        ("KO", "m3.dat", "/x/m3.dat", "m3c0", "m3", "0:3", "male", "KO, adult", 16.0, 3.0, "rigB"),
        ("KO", "m4.dat", "/x/m4.dat", "m4c0", "m4", "0:3", "female", "KO, adult", 16.5, 3.2, "rigB"),
    ]
    for r in wide_rows:
        ws.append(list(r))

    # One per-event sheet (name truncated-style suffix _events).
    ws = wb.create_sheet("amplitude_pA_events")
    ws.append(["cell_id", "animal_id", "group", "file_name", "event_index", "value"])
    for animal, grp in (("m1", "WT"), ("m3", "KO")):
        for i in range(5):
            base = 8.0 if grp == "WT" else 14.0
            ws.append([f"{animal}c0", animal, grp, f"{animal}.dat", i, base + i * 0.5])

    wb.save(path)
    return path


def build_nednet_results(path: Path) -> Path:
    """results_graph_data.xlsx: PerAnimal, Daily_counts_byGroup (two-row header),
    ISI_distribution, ISI_CDF."""
    wb = Workbook()

    # PerAnimal (dict table) with an excluded animal.
    ws = wb.active
    ws.title = "PerAnimal"
    ws.append(["animal", "group", "cohort", "rec_hours", "rec_days", "events",
               "convulsive", "non_convulsive", "rate_per_h", "conv_rate_per_h",
               "nonconv_rate_per_h", "mean_duration_s", "excluded"])
    per_animal = [
        ("a1", "sham", "c1", 100, 5, 10, 2, 8, 0.10, 0.02, 0.08, 12.3, False),
        ("a2", "sham", "c1", 100, 5, 12, 3, 9, 0.12, 0.03, 0.09, 11.8, False),
        ("a3", "kindled", "c1", 100, 5, 40, 20, 20, 0.40, 0.20, 0.20, 22.1, False),
        ("a4", "kindled", "c1", 100, 5, 38, 18, 20, 0.38, 0.18, 0.20, 20.7, False),
        ("a5", "kindled", "c1", 100, 5, 0, 0, 0, 0.00, 0.00, 0.00, 0.0, True),  # excluded
    ]
    for r in per_animal:
        ws.append(list(r))

    # Daily_counts_byGroup (two-row header via _wide):
    #   row0: ['', group, group, ...]  row1: ['recording_day', animal, ...]
    ws = wb.create_sheet("Daily_counts_byGroup")
    animals = ["a1", "a2", "a3", "a4"]
    groups = ["sham", "sham", "kindled", "kindled"]
    ws.append([""] + groups)
    ws.append(["recording_day"] + animals)
    for day in (1, 2, 3):
        ws.append([day] + [day * (1 if g == "sham" else 4) for g in groups])

    # ISI_distribution (pre-binned, one column per group).
    ws = wb.create_sheet("ISI_distribution")
    ws.append(["ISI_low_s", "ISI_high_s", "ISI_center_s", "sham", "kindled"])
    for i in range(4):
        lo, hi = 0.1 * (i + 1), 0.1 * (i + 2)
        ws.append([lo, hi, (lo + hi) / 2, 0.5 - 0.1 * i, 0.3 + 0.05 * i])

    # ISI_CDF.
    ws = wb.create_sheet("ISI_CDF")
    ws.append(["ISI_s", "sham", "kindled"])
    for j in range(5):
        ws.append([0.1 * j, min(1.0, 0.2 * j), min(1.0, 0.25 * j)])

    wb.save(path)
    return path
