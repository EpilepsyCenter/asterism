"""Reading spreadsheets into raw rows: CSV / TSV / text, Excel (.xlsx, .xlsm) and old
Excel (.xls). Every sheet becomes a list of rows of plain values (str, float, int,
datetime or None); numbers written as text (including decimal commas) become numbers."""

from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from pathlib import Path

SPREADSHEET_EXTS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}
_NUM = re.compile(r"^[+-]?(\d+([.,]\d*)?|[.,]\d+)([eE][+-]?\d+)?$")


def number(v):
    """A float for numeric cells and numeric text ("1.5", "1,5", " 2e-3 "), else None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v else None          # NaN → None
    if isinstance(v, str):
        s = v.strip().replace("−", "-")
        if _NUM.match(s):
            if s.count(",") == 1 and "." not in s:
                s = s.replace(",", ".")
            try:
                return float(s)
            except ValueError:
                return None
    return None


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat(sep=" ") if isinstance(v, _dt.datetime) and (v.hour or v.minute or v.second) \
            else v.strftime("%Y-%m-%d")
    return str(v).strip()


def _clean(v):
    if isinstance(v, str):
        n = number(v)
        return n if n is not None else (v.strip() or None)
    return v


def read_text(data: str) -> list[list]:
    """CSV/TSV text: the delimiter is sniffed (comma, semicolon, tab)."""
    delim = _delimiter(data[:50000])
    rows = [[_clean(c) for c in r] for r in csv.reader(io.StringIO(data), delimiter=delim)]
    return _trim(rows)


def _delimiter(sample: str) -> str:
    """The delimiter most lines agree on (tab, semicolon or comma). Decimal commas and a
    preamble don't fool it: what counts is how many lines share the same number of fields."""
    lines = [ln for ln in sample.splitlines() if ln.strip()][:200]
    best, best_score = ",", -1.0
    for d in ("\t", ";", ","):
        counts = [ln.count(d) for ln in lines]
        nz = [c for c in counts if c > 0]
        if not nz:
            continue
        mode = max(set(nz), key=nz.count)
        score = nz.count(mode) + (0.5 if d != "," else 0.0)
        if score > best_score:
            best, best_score = d, score
    return best


def _trim(rows: list[list]) -> list[list]:
    """Drop trailing empty rows and columns."""
    while rows and all(is_blank(c) for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for j in range(len(r) - 1, -1, -1):
            if not is_blank(r[j]):
                width = max(width, j + 1)
                break
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]


def read_file(path: str | Path) -> dict[str, list[list]]:
    """{sheet name: rows}. CSV/TSV/text files have one sheet named after the file."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = {ws.title: _trim([[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)])
               for ws in wb.worksheets}
        wb.close()
        return out
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path))
        out = {}
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                        except Exception:
                            row.append(cell.value)
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(_clean(cell.value))
                rows.append(row)
            out[sh.name] = _trim(rows)
        return out
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            data = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return {path.stem: read_text(data)}


def guess_header_row(rows: list[list], max_scan: int = 20) -> int:
    """The row holding column names: the first mostly-text row followed by data."""
    best, best_score = 0, -1.0
    for i, row in enumerate(rows[:max_scan]):
        cells = [c for c in row if not is_blank(c)]
        if len(cells) < 1:
            continue
        frac_text = sum(1 for c in cells if number(c) is None) / len(cells)
        below = rows[i + 1] if i + 1 < len(rows) else []
        nb = [c for c in below if not is_blank(c)]
        num_below = (sum(1 for c in nb if number(c) is not None) / len(nb)) if nb else 0.0
        score = frac_text + 0.25 * num_below + 0.6 * min(len(cells), 10) / 10 - 0.03 * i
        if score > best_score:
            best, best_score = i, score
    return best


def read_files(paths) -> dict[str, list[list]]:
    """Several files (or an export folder of CSVs, as QuantStudio sometimes writes) as one
    workbook: each file's sheets, named after the file when needed."""
    out: dict[str, list[list]] = {}
    files = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            files += sorted(f for f in p.iterdir() if f.suffix.lower() in SPREADSHEET_EXTS and not f.name.startswith("."))
        else:
            files.append(p)
    for f in files:
        for name, rows in read_file(f).items():
            key = name if name not in out else f"{f.stem} — {name}"
            out[key] = rows
    return out


def to_jsonable(sheets: dict) -> dict:
    """Rows as JSON (dates as {"$date": iso}); kept in a project for changing import settings."""
    def cell(v):
        if isinstance(v, (_dt.datetime, _dt.date)):
            return {"$date": v.isoformat()}
        return v
    return {name: [[cell(c) for c in r] for r in rows] for name, rows in sheets.items()}


def from_jsonable(data: dict) -> dict:
    def cell(v):
        if isinstance(v, dict) and "$date" in v:
            try:
                return _dt.datetime.fromisoformat(v["$date"])
            except ValueError:
                return v["$date"]
        return v
    return {name: [[cell(c) for c in r] for r in rows] for name, rows in data.items()}
