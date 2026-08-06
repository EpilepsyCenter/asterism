"""Source recipes for TRACER and NED-Net (SPEC §6).

Derived directly from the frozen apps' export writers (TRACER
`backend/analysis/cohort_export.py`, NED-Net `dash_app/pages/results.py`), so these
match the real files by construction. Sheets are matched by **header signature**, not
by name, because both writers truncate sheet names to Excel's 31-char limit.

Each recipe returns a canonical long `Dataset` (or, for the pre-binned distribution
sheets, a tidy distribution frame that feeds the frequency-distribution / KS figures).
Nothing is inferred silently — the returned objects record what was mapped, and the
callers surface it.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from ..contract import (METRIC, REPLICATE, SOURCE_FILE, SUBJECT, VALUE, Dataset)
from .reshape import (assign_tags_to_factors, melt_to_long,
                     two_row_header_to_long)
from .roles import Role

# --- header signatures (from the writers) ----------------------------------
TRACER_LONG_HEADER = ["cell_id", "animal_id", "group", "file_name", "metric", "value"]
TRACER_EVENT_HEADER = ["cell_id", "animal_id", "group", "file_name",
                       "event_index", "value"]
TRACER_WIDE_PREFIX = ["group", "file_name", "file_path", "cell_id", "animal_id",
                      "series_key (0-indexed)", "series_specific_tags", "group_tags"]

NEDNET_TWO_ROW_XLABELS = {"recording_day", "hour_of_day"}
NEDNET_ISI_DIST_FIRST = "ISI_low_s"
NEDNET_ISI_CDF_FIRST = "ISI_s"
# PerAnimal / Duration_perAnimal numeric metric columns to melt (the rest are ids).
NEDNET_PERANIMAL_METRICS = ["rate_per_h", "conv_rate_per_h", "nonconv_rate_per_h",
                            "mean_duration_s"]
NEDNET_DURATION_METRICS = ["mean_dur_conv_s", "mean_dur_nonconv_s",
                           "mean_dur_all_s", "median_dur_all_s"]


# --- raw workbook access ----------------------------------------------------

def load_sheets_raw(path: Path | str) -> dict[str, list[list]]:
    """Every sheet as a list of raw rows (values only). Needed because the
    two-row-header sheets can't be read with a single pandas header."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[list]] = {}
    for ws in wb.worksheets:
        out[ws.title] = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def header_of(rows: list[list]) -> list[str]:
    if not rows:
        return []
    return [("" if c is None else str(c)) for c in rows[0]]


def _frame_from_rows(rows: list[list]) -> pd.DataFrame:
    hdr = header_of(rows)
    body = rows[1:]
    return pd.DataFrame(body, columns=hdr)


# --- classification ---------------------------------------------------------

@dataclass(frozen=True)
class SheetPlan:
    sheet: str
    source: str            # 'tracer' | 'nednet'
    kind: str              # 'cells_long' | 'cells_wide' | 'events' | 'per_animal'
                           # | 'two_row' | 'isi_dist' | 'isi_cdf'
    metric: str | None = None


def classify_sheets(sheets: dict[str, list[list]]) -> list[SheetPlan]:
    """Classify every sheet by header signature (SPEC §6)."""
    plans: list[SheetPlan] = []
    for name, rows in sheets.items():
        hdr = header_of(rows)
        low = [h.lower() for h in hdr]
        if hdr[:6] == TRACER_LONG_HEADER:
            plans.append(SheetPlan(name, "tracer", "cells_long"))
        elif hdr[:6] == TRACER_EVENT_HEADER:
            # metric name is the sheet title minus the writer's '_events' suffix
            metric = name[:-7] if name.endswith("_events") else name
            plans.append(SheetPlan(name, "tracer", "events", metric=metric))
        elif hdr[:len(TRACER_WIDE_PREFIX)] == TRACER_WIDE_PREFIX:
            plans.append(SheetPlan(name, "tracer", "cells_wide"))
        elif hdr and hdr[0] == NEDNET_ISI_DIST_FIRST:
            plans.append(SheetPlan(name, "nednet", "isi_dist"))
        elif hdr and hdr[0] == NEDNET_ISI_CDF_FIRST:
            plans.append(SheetPlan(name, "nednet", "isi_cdf"))
        elif "animal" in low and any(m in low for m in NEDNET_PERANIMAL_METRICS
                                     + NEDNET_DURATION_METRICS):
            plans.append(SheetPlan(name, "nednet", "per_animal"))
        elif _is_two_row_header(rows):
            metric = name  # sheet title carries the meaning (Daily_counts_byGroup…)
            plans.append(SheetPlan(name, "nednet", "two_row", metric=metric))
    return plans


def _is_two_row_header(rows: list[list]) -> bool:
    if len(rows) < 3 or not rows[1]:
        return False
    first_data_label = rows[1][0]
    return (str(first_data_label) in NEDNET_TWO_ROW_XLABELS
            and (rows[0][0] in (None, "")))


# --- TRACER recipes ---------------------------------------------------------

def import_tracer_cells_long(rows: list[list]) -> Dataset:
    """`Cells (long)` → canonical long, no reshaping (SPEC §6)."""
    df = _frame_from_rows(rows)
    mapping = {
        "cell_id": Role.REPLICATE, "animal_id": Role.SUBJECT,
        "group": Role.GROUP, "file_name": Role.SOURCE_FILE,
        "metric": Role.METRIC, "value": Role.VALUE,
    }
    from .reshape import apply_roles
    return apply_roles(df, mapping)


def import_tracer_events(rows: list[list], metric: str | None = None) -> Dataset:
    """A per-event sheet → long, one row per event (for frequency dists / KS)."""
    df = _frame_from_rows(rows)
    if metric is not None:
        df = df.assign(**{METRIC: metric})
    mapping = {
        "cell_id": Role.REPLICATE, "animal_id": Role.SUBJECT,
        "group": Role.GROUP, "file_name": Role.SOURCE_FILE, "value": Role.VALUE,
    }
    if METRIC in df.columns:
        mapping[METRIC] = Role.METRIC
    from .reshape import apply_roles
    # event_index stays as an extra column (finer than replicate); not a role.
    return apply_roles(df, mapping)


def import_tracer_cells_wide(
    rows: list[list],
    *,
    metrics: list[str] | None = None,
    tag_assignment: dict[str, str] | None = None,
) -> Dataset:
    """`Cells (wide)` → long, melting the chosen scalar metric columns and
    optionally splitting the joined `group_tags` into factor columns (SPEC §5.3, §6).
    """
    df = _frame_from_rows(rows)
    id_cols = ["cell_id", "animal_id", "group", "file_name"]
    # tag map first, so the new factor columns ride along as id_vars
    factor_cols: list[str] = ["group"]
    if tag_assignment:
        df = assign_tags_to_factors(df, "group_tags", tag_assignment)
        factor_cols += sorted(set(tag_assignment.values()))

    if metrics is None:
        # every column after the fixed prefix that isn't meta.* is a scalar metric
        after = list(df.columns)[len(TRACER_WIDE_PREFIX):]
        metrics = [c for c in after if not str(c).startswith("meta.")
                   and c not in factor_cols and c != "tags_unassigned"]

    keep = [c for c in ["animal_id", "cell_id", "file_name", *factor_cols]
            if c in df.columns]
    long = melt_to_long(df, id_vars=keep, value_vars=metrics)

    mapping = {
        "animal_id": Role.SUBJECT, "cell_id": Role.REPLICATE,
        "file_name": Role.SOURCE_FILE, METRIC: Role.METRIC, VALUE: Role.VALUE,
    }
    for f in factor_cols:
        if f in long.columns:
            mapping[f] = Role.GROUP
    from .reshape import apply_roles
    return apply_roles(long, mapping)


# --- NED-Net recipes --------------------------------------------------------

@dataclass
class ImportResult:
    """A dataset plus anything the importer must surface loudly (SPEC §2, §6)."""

    dataset: Dataset
    excluded_subjects: tuple[str, ...] = ()
    notes: tuple[str, ...] = ()


def import_nednet_per_animal(
    rows: list[list],
    *,
    metrics: list[str] | None = None,
) -> ImportResult:
    """`PerAnimal` → long. Melts the rate/duration metric columns; `animal`→subject,
    `group`→group. Respects the `excluded` column: excluded animals are reported and
    dropped from the dataset but the count is surfaced, never silent (SPEC §6)."""
    df = _frame_from_rows(rows)
    excluded_subjects: tuple[str, ...] = ()
    if "excluded" in df.columns:
        excl_mask = df["excluded"].map(_truthy)
        excluded_subjects = tuple(
            str(a) for a in df.loc[excl_mask, "animal"].dropna().unique()
        ) if "animal" in df.columns else ()
        df = df.loc[~excl_mask].copy()

    if metrics is None:
        metrics = [m for m in NEDNET_PERANIMAL_METRICS + NEDNET_DURATION_METRICS
                   if m in df.columns]
    id_cols = [c for c in ["animal", "group", "cohort"] if c in df.columns]
    long = melt_to_long(df, id_vars=id_cols, value_vars=metrics)

    mapping = {"animal": Role.SUBJECT, "group": Role.GROUP, METRIC: Role.METRIC,
               VALUE: Role.VALUE}
    if "cohort" in long.columns:
        mapping["cohort"] = Role.BATCH
    from .reshape import apply_roles
    ds = apply_roles(long, mapping)
    notes = ()
    if excluded_subjects:
        notes = (f"{len(excluded_subjects)} animal(s) excluded per source file: "
                 f"{', '.join(excluded_subjects)}",)
    return ImportResult(ds, excluded_subjects, notes)


def import_nednet_two_row(rows: list[list], metric: str) -> Dataset:
    """A `Daily_*`/`Circadian_byGroup` two-row-header block → long (SPEC §6)."""
    long = two_row_header_to_long(rows, metric=metric)
    mapping = {SUBJECT: Role.SUBJECT, "group": Role.GROUP, METRIC: Role.METRIC,
               VALUE: Role.VALUE}
    from .reshape import apply_roles
    # the x column (recording_day / hour_of_day) is a numeric predictor (§8)
    x_col = next((c for c in long.columns if c in NEDNET_TWO_ROW_XLABELS), None)
    return apply_roles(long, mapping, x=x_col)


def parse_nednet_isi(rows: list[list], *, kind: str) -> pd.DataFrame:
    """Pre-binned ISI distribution / CDF → tidy long distribution frame.

    These are group-pooled (no subject), so they feed the frequency-distribution and
    KS figures directly rather than the subject-based stats path (SPEC §6).
    Returns columns [x, group, value] where x is the bin centre / edge.
    """
    df = _frame_from_rows(rows)
    if kind == "isi_dist":
        x_col = "ISI_center_s"
        group_cols = [c for c in df.columns
                      if c not in ("ISI_low_s", "ISI_high_s", "ISI_center_s")]
    else:  # isi_cdf
        x_col = "ISI_s"
        group_cols = [c for c in df.columns if c != "ISI_s"]
    long = df.melt(id_vars=[x_col], value_vars=group_cols,
                   var_name="group", value_name=VALUE)
    long = long.rename(columns={x_col: "x"})
    for c in ("x", VALUE):
        long[c] = pd.to_numeric(long[c], errors="coerce")
    return long


def _truthy(v) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().lower() in {"true", "1", "yes", "y", "excluded"}
